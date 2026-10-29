"""Diapositivas del examen parcial: portada + 7 de contenido.

Uso:
    python EP/slides/construir_slides.py

Sale Examen_Parcial_CC0F4.pptx y .pdf (el pdf con PowerPoint, que no se
cierra si ya tenia otras presentaciones abiertas). Las cifras se leen de
EP/results/.
"""
import sys
from pathlib import Path

import pandas as pd
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

aqui = Path(__file__).resolve().parent
ep = aqui.parent
res = ep / "results"
sys.path.insert(0, str(ep.parent / "pc2" / "slides"))
from comun import exportar, nueva_presentacion  # noqa: E402

prs = nueva_presentacion()
NEGRO = RGBColor(0, 0, 0)


def slide_titulo(titulo):
    s = prs.slides.add_slide(prs.slide_layouts[5])
    t = s.shapes.title
    t.left, t.top, t.width, t.height = Inches(0.6), Inches(0.3), Inches(12.1), Inches(0.9)
    t.text_frame.text = titulo
    t.text_frame.paragraphs[0].runs[0].font.size = Pt(32)
    t.text_frame.paragraphs[0].alignment = 1
    return s


def bullets(s, items, x=0.6, y=1.4, w=12.1, h=5.6, size=20):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        nivel = 1 if it.startswith("  ") else 0
        p.text = ("- " if nivel else "• ") + it.strip()
        p.level = nivel
        p.space_after = Pt(6)
        for r in p.runs:
            r.font.size = Pt(size - 3 * nivel)
    return tb


def tabla(s, datos, x, y, anchos, alto=0.36, size=13):
    gf = s.shapes.add_table(len(datos), len(datos[0]), Inches(x), Inches(y), Inches(sum(anchos)),
                            Inches(alto * len(datos)))
    t = gf.table
    for j, a in enumerate(anchos):
        t.columns[j].width = Inches(a)
    for i, fila in enumerate(datos):
        t.rows[i].height = Inches(alto)
        for j, v in enumerate(fila):
            c = t.cell(i, j)
            c.text = str(v)
            c.margin_top = c.margin_bottom = Inches(0.03)
            for p in c.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(size)
    return t


def caja(s, x, y, w, h, texto, size=14):
    b = s.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    b.fill.background()
    b.line.color.rgb = NEGRO
    b.line.width = Pt(1.25)
    b.shadow.inherit = False
    tf = b.text_frame
    tf.word_wrap = True
    for i, linea in enumerate(texto.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = linea
        p.alignment = 2
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.color.rgb = NEGRO
    return b


def flecha(s, x1, y1, x2, y2):
    from lxml import etree
    c = s.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = NEGRO
    c.line.width = Pt(1.5)
    e = etree.SubElement(c.line._get_or_add_ln(), "{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd")
    e.set("type", "triangle")


def notas(s, texto):
    s.notes_slide.notes_text_frame.text = texto


def n(x, d=2):
    return f"{x:.{d}f}".replace(".", ",")


r1 = pd.read_csv(res / "exp1" / "resumen.csv", index_col=0)
lat = pd.read_csv(res / "exp1" / "latencia.csv", index_col=0)
r2 = pd.read_csv(res / "exp2" / "resumen.csv", index_col=0)
r3 = pd.read_csv(res / "exp3" / "resumen.csv", index_col=0)

# portada
s = prs.slides.add_slide(prs.slide_layouts[0])
s.shapes.title.text = "Examen Parcial - Tópicos de Ciencia de la Computación 1"
s.placeholders[1].text = "Recuperación, RAG y contratos de herramientas sobre registros OCDS del SEACE"
s.shapes.title.left, s.shapes.title.top, s.shapes.title.width, s.shapes.title.height = Inches(0.9), Inches(2.0), Inches(11.5), Inches(1.8)
s.placeholders[1].left, s.placeholders[1].top, s.placeholders[1].width, s.placeholders[1].height = Inches(0.9), Inches(4.1), Inches(11.5), Inches(1.0)

# 1. problema y continuidad
s = slide_titulo("Problema y continuidad con la PC1 y la PC2")
bullets(s, [
    "Un requisito está direccionado cuando favorece a una marca, por ejemplo pidiéndola sin \"o equivalente\"",
    "  La evidencia está repartida en muchos procesos parecidos y una respuesta sin fuente no sirve",
    "PC1: Qwen2.5-0.5B clasifica un requerimiento OCDS en un JSON validado, pero no busca evidencia",
    "PC2: recuperador denso sobre 10 fichas, Recall@3 = 0,47",
    "  El chunk separaba la buena pro del ítem que la identifica, y \"o equivalente\" se perdía en el embedding",
    "EP: 55 procesos reales, 20 consultas con qrels, comparación A/B/C, RAG con citas y una herramienta con contratos",
    "  Corrección a la PC2: la unidad es un pasaje + el encabezado de su proceso (la config. de la PC2 queda como A0)",
    "Pregunta: ¿qué recuperación encuentra mejor la evidencia, cuánto cuesta, y con esa evidencia el LLM responde bien y sustentado?",
], size=20)
notas(s, "Seguimos con el mismo tema. La PC1 clasificaba un requerimiento, la PC2 buscaba evidencia con un "
         "recuperador denso. En el parcial juntamos las dos cosas y medimos cada pieza por separado.")

# 2. arquitectura y datos
s = slide_titulo("Arquitectura y datos")
y = 1.55
pasos = [("consulta", 1.3), ("e5 + FAISS\n(A)", 1.5), ("+ BM25 + RRF\n(B)", 1.6), ("+ Cross-Encoder\n(C)", 1.8),
         ("top-5 pasajes\ncon su id", 1.6), ("Qwen2.5-1.5B\ngreedy", 1.6), ("JSON: respuesta,\ncitas, abstencion", 1.9)]
x = 0.45
for i, (t, w) in enumerate(pasos):
    caja(s, x, y, w, 0.9, t, size=13)
    if i < len(pasos) - 1:
        flecha(s, x + w, y + 0.45, x + w + 0.2, y + 0.45)
    x += w + 0.2
caja(s, 0.45, 2.85, 1.3, 0.75, "ocid", size=13)
flecha(s, 1.75, 3.22, 1.95, 3.22)
caja(s, 1.95, 2.85, 3.3, 0.75, "ocds.indicadores_proceso\ncontrato entrada -> dominio -> salida", size=13)
flecha(s, 5.25, 3.22, 5.45, 3.22)
caja(s, 5.45, 2.85, 2.6, 0.75, "ToolResult / ToolError\nreintento selectivo", size=13)
tabla(s, [
    ["Datos", ""],
    ["Corpus", "55 procesos reales del SEACE (OCDS), 426 pasajes"],
    ["Pasajes", "convocatoria, ítems, cronograma, documentos, postores, resultado, contrato"],
    ["Consultas", "22: 20 con qrels (11 difíciles, distractores reales) y 2 controles sin evidencia"],
    ["qrels", "anotados leyendo los pasajes y congelados con SHA-256 antes de correr (b623b36)"],
    ["Modelos", "e5-small, mMiniLM Cross-Encoder, Qwen2.5-1.5B; revisiones fijadas, CPU"],
], 0.45, 4.0, [1.6, 10.5], alto=0.42, size=14)
notas(s, "Arriba el flujo RAG: la consulta pasa por las tres condiciones del experimento 1 y los 5 primeros "
         "pasajes de C van al LLM con su id. Abajo la herramienta, que es independiente: no hay agente que las conecte.")

# 3. trazabilidad
s = slide_titulo("Trazabilidad de las semanas 1 a 7")
tabla(s, [
    ["Semana", "Recurso", "Dónde se aplica", "Evidencia / conclusión"],
    ["1", "Cuaderno 1 §4-6, Lectura 2 §2-3", "recuperacion.py: bi-encoder e5 vs Cross-Encoder; Qwen decoder-only", "C gana MRR +0,16 sobre A, pero cuesta ~48 veces más"],
    ["2", "Cuaderno 2 §8-12, Lectura 3 §8, Lab 2 §10", "exp. 2: greedy; KV cache con GQA (2 cabezas KV)", "271 vs 943 tokens; 20 s vs 35 s; ~56 KiB de cache por token"],
    ["3", "Cuaderno 3 §3, §9-10", "rag.parsear: JSON {respuesta, citas, abstencion}", "24/24 JSON válidos, 4 incorrectas: formato ≠ corrección"],
    ["4", "Cuaderno 4 §6-7, Lab 4 §3", "A0 (chunks PC2) vs A (pasaje + encabezado), IndexFlatIP", "A − A0: MRR +0,18 [0,05; 0,32]"],
    ["5", "Cuaderno 5 §3, §5-6, §10", "B = BM25 + RRF k = 60; C = Cross-Encoder", "Q06: rank 4 -> 1 con BM25; B ≈ A en promedio"],
    ["6", "Cuaderno 6 §2-5, Lectura 6 §9-14", "herramienta.py: contratos, ToolError, políticas A-D", "A 6/16, D 16/16 sin reintentos inútiles"],
    ["7", "Cuaderno 7 §3, §5.1, §6-9, Lectura 7 §19", "metricas.py, auditoría de afirmaciones", "solo C − A no cruza cero; citas auto ≠ manual"],
], 0.45, 1.4, [0.9, 3.3, 4.3, 3.9], alto=0.62, size=13)
notas(s, "Cada fila cita la sección del material del curso y el archivo del EP donde se aplica; "
         "el detalle, con las funciones de cada cuaderno, está en la sección 5 del README.")

# 4. experimento 1
s = slide_titulo("Experimento 1: recuperación densa, híbrida y con reranking")
s.shapes.add_picture(str(res / "figs" / "exp1_metricas.png"), Inches(0.3), Inches(1.25), Inches(7.4))
tabla(s, [
    ["", "Latencia media", "p95"],
    *[[c, f"{lat.loc[c, 'total_ms']:.0f} ms", f"{lat.loc[c, 'total_p95_ms']:.0f} ms"] for c in ["A0", "A", "B", "C"]],
], 7.95, 1.35, [0.9, 1.9, 1.3], alto=0.38, size=13)
tabla(s, [
    ["Bootstrap (MRR)", "Δ e IC 95 %"],
    ["B − A", "+0,02 [−0,10; 0,15]"],
    ["C − B", "+0,14 [−0,01; 0,30]"],
    ["C − A", "+0,16 [0,02; 0,31]"],
], 7.95, 3.45, [1.9, 2.2], alto=0.38, size=13)
bullets(s, [
    "Q06 (código 006R01830): rank 4 en A, rank 1 con BM25",
    "Q20, Q12: solo el Cross-Encoder pone la buena pro / los postores en el rank 1",
    "Q03: el reranker baja el ítem correcto al rank 9 y prefiere el distractor (Adobe sin equivalente)",
    "Q13 (> 30 % sobre el valor referencial) falla en todas: hay que comparar números",
    "Negativo: Recall@5 de C (0,90) < B (0,94); B − A y C − B cruzan cero",
], x=0.45, y=4.85, w=12.4, h=2.5, size=16)
notas(s, "A frente a B cambia dos cosas a la vez, BM25 y la fusión. B frente a C solo reordena los mismos 20 "
         "candidatos. El reranker ayuda arriba de la lista pero cuesta 1 segundo por consulta. Con 20 consultas "
         "solo C contra A es distinguible de cero.")

# 5. experimento 2
s = slide_titulo("Experimento 2: closed-book vs RAG")
s.shapes.add_picture(str(res / "figs" / "exp2_auditoria.png"), Inches(0.3), Inches(1.25), Inches(7.2))
tabla(s, [
    ["Auditoría manual", "Closed-book", "RAG"],
    ["Correctas / parciales", "0 / 0", f"{r2.loc['rag_C', 'correctas']} / {r2.loc['rag_C', 'parciales']}"],
    ["Abstención en controles", r2.loc["closed_book", "abstencion_correcta_en_controles"],
     r2.loc["rag_C", "abstencion_correcta_en_controles"]],
    ["Faithfulness", "0/1", "16/19"],
    ["Citas que sustentan", "-", r2.loc["rag_C", "citas_que_sustentan_manual"]],
    ["Citas en qrels (auto)", "-", r2.loc["rag_C", "citas_en_qrels_auto"]],
], 7.8, 1.35, [2.6, 1.4, 1.2], alto=0.4, size=13)
bullets(s, [
    "Mismo Qwen2.5-1.5B, 12 preguntas fijadas y greedy; solo cambia el contexto (vacío vs top-5 de C)",
    "Q21 (sanción, sin evidencia): \"el Tribunal impuso una sanción de S/ 184,500.00\" citando la buena pro E01-P06",
    "Q22: inventa un plazo de \"30 días\"; Q15: dice que no aceptaba equivalentes y luego que sí",
    "Q12: respuesta correcta, pero cita E50-P01 en vez de E50-P05 (postores)",
    "Q04: la cita está en qrels pero no sustenta lo que dice: auto 12/25 y manual 12/25 no son las mismas citas",
], x=0.45, y=4.35, w=12.4, h=3.0, size=16)
notas(s, "Sin evidencia el modelo casi siempre se abstiene. Con evidencia responde bien 8 de 10, pero pierde la "
         "abstención: el recuperador siempre devuelve 5 pasajes y el modelo los usa aunque no respondan la pregunta.")

# 6. experimento 3
s = slide_titulo("Experimento 3: herramienta ocds.indicadores_proceso")
bullets(s, [
    "Entrada: ocid (patrón SEACE) y tender_item_id (entero estricto); extra=forbid",
    "Dominio: el ocid existe, el ítem es del proceso, hay que elegir ítem si hay varios",
    "Salida: único postor, adjudicado / valor ref., \"o equivalente\" en objeto vs ítem; tipos y coherencia",
    "Se reintenta solo transient y timeout (solo lectura, 3 intentos, 100 ms por intento, deadline 300 ms)",
    "  No: validation, domain, output_contract, execution (reintentar da el mismo error)",
], x=0.45, y=1.3, w=12.4, h=2.6, size=17)
tabla(s, [
    ["Política (16 casos)", "Cumplen", "Entradas inválidas ejecutadas", "Salidas inválidas aceptadas", "Intentos", "Reintentos inútiles"],
    *[[nombre, f"{int(r3.loc[p, 'cumplen_especificacion'])}/16", int(r3.loc[p, "entradas_invalidas_ejecutadas"]),
       int(r3.loc[p, "salidas_invalidas_aceptadas"]), int(r3.loc[p, "intentos_totales"]), int(r3.loc[p, "reintentos_inutiles"])]
      for p, nombre in [("A", "A sin contratos"), ("B", "B contratos"), ("C", "C reintenta todo"), ("D", "D reintento selectivo")]],
], 0.45, 4.05, [2.9, 1.2, 2.6, 2.5, 1.2, 1.9], alto=0.4, size=14)
bullets(s, [
    "A acepta un argumento inventado y devuelve ok; también deja pasar n_postores = −1",
    "22 pruebas con pytest: válida, inválida, dominio, transitorio, timeout, error interno, salida rota",
], x=0.45, y=6.25, w=12.4, h=1.0, size=15)
notas(s, "La herramienta no decide si hay direccionamiento: da hechos del registro para contrastar al LLM. "
         "En Q15 el LLM se contradice y la herramienta devuelve la alerta equivalente_solo_en_objeto.")

# 7. resultados, errores y conclusiones
s = slide_titulo("Resultados, errores y conclusiones")
bullets(s, [
    "La unidad de recuperación importó más que el método: pasaje + encabezado subió el MRR en 0,18 frente a la PC2",
    "BM25 + RRF no mejora en promedio; ayuda con códigos exactos y empeora cuando la fusión adelanta un distractor",
    "El Cross-Encoder mejora el inicio del ranking (Recall@1 0,41 -> 0,61) a ~1 s por consulta; contra B no es concluyente",
    "RAG: 8 correctas y 1 parcial de 10, pero 0/2 abstenciones y 13 de 25 citas no sustentan nada",
    "  Una respuesta citada no es una respuesta verificada",
    "La herramienta con contratos evita ejecuciones inválidas silenciosas; reintentar todo gasta sin ganar",
    "Limitaciones: 20 consultas, qrels y auditoría nuestros, sin las bases en PDF, un solo LLM pequeño, fallos simulados",
], size=19)
notas(s, "Ninguna conclusión se generaliza fuera de este benchmark. Lo que queda para la semana 9 en adelante es "
         "un recuperador evaluado, un generador con sus fallos documentados y una herramienta con contratos.")

salida = aqui / "Examen_Parcial_CC0F4.pptx"
prs.save(salida)
exportar(salida, aqui / "Examen_Parcial_CC0F4.pdf", aqui / "_png")
print("listo", salida, "| diapositivas:", len(prs.slides))
